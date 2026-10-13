"""Supplementary Table S2: Bio-Three exposure within Site B.

Bio-Three was used at Site B until it was withdrawn from the unit's formulary on
10 July 2024, after which no infant received it.  Exposure is therefore defined from
the calendar date of each sample (date of birth + day of life): a Site B sample is
exposed if it was obtained before the withdrawal date.  An infant is counted as
exposed if any of its samples was exposed.

Requires meta.pkl, genus.pkl and cst_v2.pkl, and the clinical workbook (for date of
birth and the infant characteristics).  Set BGM_UPLOADS if the workbook is not in
/mnt/user-data/uploads/.
"""
import os
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from scipy import stats
from statsmodels.stats.multitest import multipletests

UP = os.environ.get("BGM_UPLOADS", "/mnt/user-data/uploads/")
WITHDRAWAL = pd.Timestamp("2024-07-10")
N_MC, SEED = 200_000, 20260726
NAME = {7: "C1", 1: "C2", 4: "C3", 2: "C4", 0: "C5", 5: "C6", 6: "C7", 3: "C8"}
TYPES = [f"C{i}" for i in range(1, 9)]

ws = load_workbook(UP + "BGM研究ノート学会短縮版_v2.xlsx", data_only=True)["シート1"]
rows = [list(r) for r in ws.iter_rows(values_only=True)]
cases = [str(x).strip() for x in rows[1][1:] if x is not None]
F = {str(r[0]).strip(): r[1:1 + len(cases)] for r in rows[2:] if r[0] is not None}
info = pd.DataFrame({k: F[k] for k in ["出生日", "Sex(male)", "CAM", "Bw z-score(T0)", "Bw z-score(T5)"]
                     + [f"Nutrition(T{t})" for t in range(1, 6)] + [f"ABx(T{t})" for t in range(1, 6)]},
                    index=cases)
birth = pd.to_datetime(info["出生日"].astype(str))

meta = pd.read_pickle("meta.pkl"); g = pd.read_pickle("genus.pkl"); cst = pd.read_pickle("cst_v2.pkl")
s = meta[meta.siid.isin(set(g.siid))].merge(cst, on="siid")
s["P"] = s.cluster.map(NAME)
s = s[s.site == "B"].copy()
s["date"] = s.case.map(birth) + pd.to_timedelta(s.dol, unit="D")
s["exposed"] = s.date < WITHDRAWAL

inf = s.groupby("case").agg(exposed=("exposed", "max"), n=("siid", "size"),
                            first=("exposed", "first"), all_=("exposed", "min"))
ei, ui = inf.index[inf.exposed], inf.index[~inf.exposed]
ga = meta.groupby("case").ga.first()


def num(col, idx):
    return pd.to_numeric(info.loc[idx, col], errors="coerce").dropna()


def med(v, d=1):
    return f"{v.median():.{d}f} ({v.quantile(.25):.{d}f} to {v.quantile(.75):.{d}f})"


def mw(a, b, **kw):
    return stats.mannwhitneyu(a, b, alternative="two-sided", **kw).pvalue


def fisher(a, na, b, nb):
    return stats.fisher_exact([[a, na - a], [b, nb - b]]).pvalue


out = []
def row(label, ex, un, p):
    out.append(dict(characteristic=label, exposed=ex, not_exposed=un, P=p))

row("Infants, n", len(ei), len(ui), None)
row("  exposed throughout / until withdrawal", int(inf.loc[ei, "all_"].sum()),
    int((~inf.loc[ei, "all_"]).sum()), None)
row("Gestational age, weeks", med(ga[ei]), med(ga[ui]), mw(ga[ei], ga[ui]))
a, b = num("Bw z-score(T0)", ei), num("Bw z-score(T0)", ui)
row("Birth-weight z-score at birth", med(a), med(b), mw(a, b))
a, b = num("Bw z-score(T5)", ei), num("Bw z-score(T5)", ui)
row(f"Weight z-score at term-equivalent age (n = {len(a)}/{len(b)})", med(a), med(b), mw(a, b))

def binary(label, f):
    x = sum(f(c) for c in ei); y = sum(f(c) for c in ui)
    row(label, f"{x} ({x/len(ei)*100:.0f})", f"{y} ({y/len(ui)*100:.0f})", fisher(x, len(ei), y, len(ui)))

binary("Male sex, n (%)", lambda c: int(float(info.loc[c, "Sex(male)"])) == 1)
binary("Chorioamnionitis, n (%)", lambda c: int(float(info.loc[c, "CAM"])) >= 1)
binary("Any antibiotic exposure, n (%)",
       lambda c: any(str(info.loc[c, f"ABx(T{t})"]) in ("E", "T") for t in range(1, 6)))
binary("Breast milk >80% of enteral intake at T1, n (%)", lambda c: str(info.loc[c, "Nutrition(T1)"]) == "B")
# periods with enteral feeding withheld (O) or not recorded (NA) are not counted against
binary("Breast milk >80% at all sampling periods, n (%)",
       lambda c: all(str(info.loc[c, f"Nutrition(T{t})"]) in ("B", "O", "NA", "None") for t in range(1, 6)))
row("Analysable samples per infant, n", med(inf.n[ei], 0), med(inf.n[ui], 0),
    mw(inf.n[ei], inf.n[ui], method="exact"))

ne, nu = int(s.exposed.sum()), int((~s.exposed).sum())
row(f"Community-type distribution (sample level; {ne} exposed and {nu} unexposed samples)", "", "", None)
ps = []
for k in TYPES:
    x = int(((s.P == k) & s.exposed).sum()); y = int(((s.P == k) & ~s.exposed).sum())
    p = fisher(x, ne, y, nu); ps.append(p)
    row(f"  {k}", f"{x} ({x/ne*100:.1f})", f"{y} ({y/nu*100:.1f})", p)
bh = multipletests(ps, method="fdr_bh")[1]

for k in ["C1", "C2", "C4"]:
    ever = s.groupby("case").P.apply(lambda v: (v == k).any())
    binary(f"Ever assigned to {k}, n (%)", lambda c, ever=ever: bool(ever[c]))
aff = s.groupby("case").ae_dol.first().notna()
binary("Intestinal complication, n (%)", lambda c: bool(aff[c]))

labels = pd.Categorical(s.P, categories=TYPES).codes
e = s.exposed.values.astype(int)

def chi2(ev):
    t = np.zeros((len(TYPES), 2)); np.add.at(t, (labels, ev), 1)
    ex = t.sum(1, keepdims=True) * t.sum(0, keepdims=True) / t.sum()
    return ((t - ex) ** 2 / ex).sum(), ex.min()

obs, min_exp = chi2(e)
rng = np.random.default_rng(SEED)
ge = sum(chi2(rng.permutation(e))[0] >= obs - 1e-9 for _ in range(N_MC))

if __name__ == "__main__":
    pd.set_option("display.width", 220); pd.set_option("display.max_colwidth", 80)
    t = pd.DataFrame(out)
    t["P"] = t.P.map(lambda p: "" if p is None or pd.isna(p) else f"{p:.4f}")
    print(t.to_string(index=False))
    print("\nBenjamini-Hochberg (sample-level types):", dict(zip(TYPES, np.round(bh, 3))))
    print(f"Overall distribution: chi2 = {obs:.1f}, df = {len(TYPES)-1}, smallest expected count = {min_exp:.2f}, "
          f"Monte Carlo exact P = {ge / N_MC:.4f} ({N_MC:,} permutations)")
