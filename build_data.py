import pandas as pd, numpy as np, re
from openpyxl import load_workbook

UP = "/mnt/user-data/uploads/"

# ---------- 1. 研究ノート（転置形式）を読む ----------
wb = load_workbook(UP + "BGM研究ノート学会短縮版_v2.xlsx", data_only=True)
ws = wb["シート1"]
rows = [[c for c in r] for r in ws.iter_rows(values_only=True)]

cases = [str(x).strip() for x in rows[1][1:] if x is not None]
n = len(cases)
print("n cases:", n)

fields = {}
for r in rows[2:]:
    key = r[0]
    if key is None:
        continue
    fields[str(key).strip()] = list(r[1:1+n])

# ---------- 2. long形式（症例×時点）に展開 ----------
recs = []
for j, case in enumerate(cases):
    for t in range(6):
        siid = fields[f"T{t}"][j]
        if siid is None or str(siid).strip().upper() == "NA":
            continue
        dol = fields[f"DOL(T{t})"][j]
        recs.append(dict(
            case=case, site=fields["Site"][j], t=t,
            siid=int(float(str(siid).strip())),
            dol=None if dol is None or str(dol).upper()=="NA" else float(dol),
            ga=fields["Gestational age"][j],
            sex_male=fields["Sex(male)"][j],
            bwz=fields["Bw z-score(T0)"][j],
            cam=fields["CAM"][j],
            nutrition=fields.get(f"Nutrition(T{t})", [None]*n)[j],
            abx=fields.get(f"ABx(T{t})", [None]*n)[j],
            ae_dol=fields["AE(DOL)"][j],
        ))
meta = pd.DataFrame(recs)

# GAの文字列（"28  5/7"）を数値化
def parse_ga(x):
    if isinstance(x, (int, float)):
        return float(x)
    s = str(x).strip()
    m = re.match(r"^(\d+)\s+(\d)/7$", s)
    if m:
        return int(m.group(1)) + int(m.group(2))/7
    return float(s)
meta["ga"] = meta["ga"].apply(parse_ga)
meta["ae_dol"] = pd.to_numeric(meta["ae_dol"], errors="coerce")
meta["affected"] = meta["ae_dol"].notna().astype(int)

print("\n=== 発症例 (AE(DOL)) ===")
print(meta[meta.affected==1].groupby(["case","site"])["ae_dol"].first())

# ---------- 3. バッチ付与 ----------
def batch(s):
    if 1 <= s <= 5:      return 1
    if 6 <= s <= 64:     return 2
    if 65 <= s <= 113:   return 3
    if 114 <= s <= 216:  return 4
    if 217 <= s <= 252:  return 5
    return np.nan
meta["batch"] = meta["siid"].apply(batch)
meta["instrument"] = np.where(meta["siid"] <= 113, "MiSeq", "i100")

# ---------- 4. 菌数データ ----------
g = pd.read_csv(UP + "bar_chart_3_genus.csv")
f = pd.read_csv(UP + "bar_chart_4_family.csv")
g["siid"] = g["Sample"].astype(int); f["siid"] = f["Sample"].astype(int)
print("\ngenus samples:", len(g), " family samples:", len(f))

gcols = [c for c in g.columns if c not in ("Sample","siid")]
fcols = [c for c in f.columns if c not in ("Sample","siid")]
g["total_all"] = g[gcols].sum(axis=1)
print("reads/sample: median %.0f  min %.0f  max %.0f" % (
    g.total_all.median(), g.total_all.min(), g.total_all.max()))

meta.to_pickle("meta.pkl"); g.to_pickle("genus.pkl"); f.to_pickle("family.pkl")

# ---------- 5. 突合 ----------
print("\n=== 突合 ===")
print("研究ノート上の検体数:", len(meta))
print("菌数データの検体数:", len(g))
print("ノートにありデータなし:", sorted(set(meta.siid) - set(g.siid)))
print("データにありノートなし:", sorted(set(g.siid) - set(meta.siid)))

print("\n=== バッチ別 検体数（解析済のみ） ===")
m2 = meta[meta.siid.isin(set(g.siid))]
print(m2.groupby("batch").size())
print("\n=== バッチ×施設 ===")
print(pd.crosstab(m2.batch, m2.site))
print("\n=== バッチ×時点 ===")
print(pd.crosstab(m2.batch, m2.t))
