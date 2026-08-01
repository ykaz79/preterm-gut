import pandas as pd, numpy as np, re
def ga(x):
    if isinstance(x,(int,float)): return float(x)
    t=str(x).strip(); m=re.match(r"^(\d+)\s+(\d)/7$",t)
    return int(m.group(1))+int(m.group(2))/7 if m else float(t)
from openpyxl import load_workbook

wb = load_workbook("/mnt/user-data/uploads/BGM研究ノート学会短縮版_v2.xlsx", data_only=True)
ws = wb["シート1"]; rows=[list(r) for r in ws.iter_rows(values_only=True)]
cases=[str(x).strip() for x in rows[1][1:] if x is not None]; n=len(cases)
F={str(r[0]).strip():list(r[1:1+n]) for r in rows[2:] if r[0] is not None}

cst = pd.read_pickle("cst.pkl"); meta = pd.read_pickle("meta.pkl"); g = pd.read_pickle("genus.pkl").set_index("siid")
f  = pd.read_pickle("family.pkl").set_index("siid")
den = g[[c for c in g.columns if c not in ("Sample","total_all")]].sum(axis=1)
K = (g["Klebsiella"]/den*100)
full = meta[meta.siid.isin(set(g.index))].merge(cst, on="siid")
full["C2"]=(full.cluster==1).astype(int); full["K"]=full.siid.map(K)

unaff_C2 = sorted(full[(full.affected==0)&(full.C2==1)].case.unique())
print("非合併症児で C2 に1回以上割り付けられた児:", len(unaff_C2), "例\n")

def longest_run(ts, flags):
    best=cur=0; prev=None
    for t,fl in sorted(zip(ts,flags)):
        if fl and (prev is None or t-prev==1): cur+=1
        elif fl: cur=1
        else: cur=0
        prev=t; best=max(best,cur)
    return best

rec=[]
for c in unaff_C2:
    j=cases.index(c); s=full[full.case==c].sort_values("t")
    c2t=[int(t) for t in s.loc[s.C2==1,"t"]]
    rec.append(dict(
        Infant=c, Site=F["Site"][j],
        GA=round(ga(F["Gestational age"][j]),1),
        Sex=("M" if str(F["Sex(male)"][j])=="1" else "F"),
        BWz_T0=F["Bw z-score(T0)"][j], BWz_T5=F["Bw z-score(T5)"][j],
        CAM=F["CAM"][j],
        Samples=len(s), C2_n=len(c2t), C2_timepoints="T"+",T".join(map(str,c2t)),
        Longest_run=longest_run(s.t.tolist(), s.C2.tolist()),
        Klebsiella_max=round(s.K.max(),1), Klebsiella_med=round(s.K.median(),1),
        Nutrition="/".join(str(F[f"Nutrition(T{t})"][j]) for t in range(1,6)),
        ABx="/".join(str(F[f"ABx(T{t})"][j]) for t in range(1,6)),
        Outcome="none"))
t = pd.DataFrame(rec).sort_values(["Site","C2_n"], ascending=[True,False])
pd.set_option("display.width",250)
print(t.to_string(index=False))

print("\n=== 比較: 合併症児5例（同一項目）===")
rec2=[]
for c in sorted(full[full.affected==1].case.unique()):
    j=cases.index(c); s=full[full.case==c].sort_values("t")
    pre=s[s.dol<s.ae_dol]
    c2t=[int(x) for x in s.loc[s.C2==1,"t"]]
    rec2.append(dict(Infant=c, Site=F["Site"][j], GA=round(ga(F["Gestational age"][j]),1),
        Sex=("M" if str(F["Sex(male)"][j])=="1" else "F"),
        BWz_T0=F["Bw z-score(T0)"][j], BWz_T5=F["Bw z-score(T5)"][j], CAM=F["CAM"][j],
        Samples=len(s), C2_n=len(c2t), C2_timepoints=("T"+",T".join(map(str,c2t))) if c2t else "-",
        Longest_run=longest_run(s.t.tolist(), s.C2.tolist()),
        Klebsiella_max=round(pre.K.max(),1) if len(pre) else np.nan,
        Klebsiella_med=round(pre.K.median(),1) if len(pre) else np.nan,
        Nutrition="/".join(str(F[f"Nutrition(T{t})"][j]) for t in range(1,6)),
        ABx="/".join(str(F[f"ABx(T{t})"][j]) for t in range(1,6)),
        Outcome=f"onset DOL {int(F['AE(DOL)'][j])}"))
print(pd.DataFrame(rec2).to_string(index=False))

print("\n=== コード出現値（凡例が必要）===")
nut=set(); abx=set()
for j in range(n):
    for tt in range(1,6):
        nut.add(str(F[f"Nutrition(T{tt})"][j])); abx.add(str(F[f"ABx(T{tt})"][j]))
print("Nutrition:", sorted(nut)); print("ABx      :", sorted(abx))
print("CAM      :", sorted({str(F['CAM'][j]) for j in range(n)}))
t.to_pickle("supp_c2.pkl")
